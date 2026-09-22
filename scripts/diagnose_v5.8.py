"""
Diagnostic: Find indices of all target paragraphs in v5.8 docx.
"""
from docx import Document

SRC = "/path/to/xelox_project/reports/manuscript_draft_v5.8_submission.docx"
doc = Document(SRC)

DISCUSSION_MARKER = "A comprehensive overfitting diagnostics summary"
TABLES2_MARKER = "Note: HR = hazard ratio; CI = confidence interval."
SUPP_MAT_MARKER = "Supplementary Materials"
TABLE_S33_MARKER = "Supplementary Table S33"

discussion_paras = []
table2_note_paras = []
supp_mat_paras = {}
table_s33_paras = []

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if DISCUSSION_MARKER in text:
        discussion_paras.append((i, text[:80]))
    if TABLES2_MARKER in text:
        table2_note_paras.append((i, text[:150]))
    if text == SUPP_MAT_MARKER:
        # Store next few paragraphs for context
        context = []
        for j in range(1, 5):
            if i + j < len(doc.paragraphs):
                context.append(doc.paragraphs[i + j].text.strip()[:80])
        supp_mat_paras[i] = context
    if TABLE_S33_MARKER in text:
        table_s33_paras.append(i)

print("=== Discussion duplicates ===")
for idx, snippet in discussion_paras:
    print(f"  Para {idx}: {snippet}")
print(f"  Count: {len(discussion_paras)}")

print("\n=== Table 2 Notes ===")
for idx, snippet in table2_note_paras:
    print(f"  Para {idx}: {snippet}")
print(f"  Count: {len(table2_note_paras)}")

print("\n=== Supplementary Materials blocks ===")
for idx, context in supp_mat_paras.items():
    print(f"  Starting Para {idx}:")
    for j, ctx in enumerate(context):
        print(f"    +{j+1}: {ctx}")
print(f"  Count: {len(supp_mat_paras)}")

print("\n=== Table S33 locations ===")
for idx in table_s33_paras:
    print(f"  Para {idx}: {doc.paragraphs[idx].text.strip()[:100]}")
print(f"  Count: {len(table_s33_paras)}")

print(f"\n=== Total paragraphs: {len(doc.paragraphs)} ===")
