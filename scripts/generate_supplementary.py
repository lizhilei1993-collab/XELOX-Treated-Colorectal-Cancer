"""Generate Supplementary Materials DOCX with actual data tables and figures."""
import os, pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PROJECT = r"/path/to/xelox_project"
OUT_PATH = os.path.join(PROJECT, "reports", "supplementary_materials.docx")
PNG_DIR = os.path.join(PROJECT, "results", "figures_png")
SCR_DIR = os.path.join(PROJECT, "results", "figures", "scrna_seq")

doc = Document()
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(10)
style.paragraph_format.line_spacing = 1.15

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
    return h

def add_para(text, bold=False, italic=False, size=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p

def add_table_from_df(df, caption, col_widths=None):
    """Add a formatted table from a pandas DataFrame."""
    add_para(caption, bold=True, size=9)
    rows, cols = df.shape
    table = doc.add_table(rows=rows+1, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header
    for j, col_name in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col_name)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(8)
    
    # Data
    for i, (_, row) in enumerate(df.iterrows()):
        for j, val in enumerate(row):
            cell = table.rows[i+1].cells[j]
            if isinstance(val, float):
                cell.text = f"{val:.3f}" if abs(val) < 100 else f"{val:.1f}"
            else:
                cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(8)
    doc.add_paragraph()

def add_image(path, caption, width_inches=5.0):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width_inches))
        add_para(caption, italic=True, size=9)
        add_para("")

# ====== TITLE ======
add_heading("Supplementary Materials", level=1)
add_para(
    "From Gene-Level Fingerprints to Pathway-Level Prognostic Modeling: "
    "An Explainable Machine Learning Framework for XELOX Chemoresistance in Colorectal Cancer",
    italic=True, size=10)
add_para("Li Z, et al. (2026)", size=10)
doc.add_paragraph()

# ====== TABLE OF CONTENTS ======
add_heading("Table of Contents", level=2)
toc_items = [
    "Table S1: 253 Meta-Analysis-Significant Genes",
    "Table S2-S5: Per-Cohort DEG Results",
    "Table S6: 10-Gene Fingerprint Details",
    "Table S7: ML Model Performance Summary",
    "Table S8: Benjamini-Hochberg Corrected Univariate Cox (44 pathways)",
    "Table S9: TGF-beta Paradox Analysis",
    "Table S10: Nomogram Sensitivity Diagnostics",
    "Table S11: Drug-Gene Interactions (Open Targets Platform)",
    "Table S12: Drug Sensitivity Comparison (PRS-High vs PRS-Low)",
    "Table S13: scRNA-seq Cell-Type Distribution (GSE132465)",
    "Figure S1-S8: Existing Supplementary Figures",
    "Figure S9: Drug Sensitivity Comparison",
    "Figure S10: scRNA-seq Validation (GSE132465)",
]
for item in toc_items:
    add_para(item, size=9)
doc.add_paragraph()

# ====== TABLE S1: 253 Genes ======
add_heading("Supplementary Tables", level=2)

# S1
meta_df = pd.read_csv(os.path.join(PROJECT, "results", "tables", "drug_prediction", "253_meta_genes.csv"))
meta_show = meta_df.head(30)[['gene_symbol', 'stouffer_z', 'stouffer_fdr', 'direction_consensus', 'n_cohorts']]
meta_show.columns = ['Gene', 'Stouffer Z', 'FDR', 'Direction', 'N Cohorts']
add_table_from_df(meta_show, f"Table S1. Meta-analysis-significant genes (Stouffer FDR < 0.05, n = 253). First 30 of 253 genes shown.")

# S11: Drug interactions
di_path = os.path.join(PROJECT, "results", "tables", "drug_prediction", "opentargets_drug_interactions.csv")
if os.path.exists(di_path):
    di_df = pd.read_csv(di_path)
    if len(di_df) > 0:
        di_show = di_df[['gene', 'drug_name', 'max_stage']].head(20)
        di_show.columns = ['Gene', 'Drug', 'Max Clinical Stage']
        add_table_from_df(di_show, "Table S11. Drug-gene interactions for XELOX resistance-associated genes (Open Targets Platform). Top 20 interactions shown.")

# S12: Drug sensitivity
ds_path = os.path.join(PROJECT, "results", "tables", "drug_prediction", "drug_sensitivity_PRShigh_vs_low.csv")
if os.path.exists(ds_path):
    ds_df = pd.read_csv(ds_path)
    ds_show = ds_df.copy()
    ds_show.columns = ['Drug', 'PRS-High Mean', 'PRS-Low Mean', 'Difference', 'p-value']
    add_table_from_df(ds_show, "Table S12. Drug sensitivity comparison between PRS-high and PRS-low groups (GSE39582, n = 164). Positive difference indicates higher resistance signature in PRS-high group.")

# S13: scRNA-seq cell types
add_table_from_df(
    pd.DataFrame({
        'Cell Type': ['T cells', 'Epithelial cells', 'B cells', 'Myeloids', 'Stromal cells', 'Mast cells', 'Total'],
        'Count': ['23,115', '18,534', '9,129', '6,762', '5,933', '187', '63,660'],
        'Percentage': ['36.3%', '29.1%', '14.3%', '10.6%', '9.3%', '0.3%', '100%']
    }),
    "Table S13. Cell-type distribution in GSE132465 CRC scRNA-seq dataset after quality control filtering."
)

# S8: Cox univariate
add_para("Table S8 (Reference). Benjamini-Hochberg corrected univariate Cox regression results for all 44 pathways. See results/tables/nomogram/ for full data file.", italic=True, size=9)
doc.add_paragraph()

# ====== SUPPLEMENTARY FIGURES ======
add_heading("Supplementary Figures", level=2)

# S9: Drug sensitivity
ds_fig = os.path.join(PROJECT, "results", "figures", "drug_prediction", "drug_sensitivity_comparison.png")
add_image(ds_fig, "Figure S9. Drug sensitivity comparison between PRS-high (n = 82) and PRS-low (n = 82) groups in GSE39582 XELOX subcohort. Bars show difference in mean resistance signature scores (positive = higher in resistant). *, p < 0.05; **, p < 0.01; ***, p < 0.001 (Mann-Whitney U test).")

# S10: scRNA-seq
sc_figs = [
    ("UMAP_cell_types.png", "Figure S10A. UMAP visualization of 63,660 CRC cells colored by 6 major cell types."),
    ("UMAP_tumor_vs_normal.png", "Figure S10B. UMAP colored by tumor vs. normal tissue origin."),
    ("resistance_genes_UMAP.png", "Figure S10C. Expression of XELOX resistance-associated genes projected onto UMAP: fingerprint score, meta-analysis up/down-regulated gene scores, and capecitabine enzyme score."),
    ("gene_scores_by_celltype.png", "Figure S10D. Mean resistance gene scores by cell type. Fingerprint genes (CSNK1G2, KAZN, KLK6, MGA, MID2) show highest expression in epithelial cells."),
    ("cape_enz_by_celltype.png", "Figure S10E. Mean capecitabine metabolism enzyme expression (CES1, CES2, DPYD, DCK, CDA) by cell type. Highest expression in epithelial and stromal cells."),
    ("tumor_vs_normal_scores.png", "Figure S10F. Boxplot comparing tumor vs. normal cells for fingerprint score and capecitabine enzyme score (both p < 0.001, Mann-Whitney U test)."),
]
for fname, caption in sc_figs:
    add_image(os.path.join(SCR_DIR, fname), caption)

# S1-S8 reference
add_para("Figures S1-S8 (Reference). See figures/ directory for the following supplementary figures:", bold=True, size=9)
for fig in ["S1: WGCNA module-trait heatmap", "S2: GDSC IC50 correlation", "S3: SHAP dependence plots",
            "S4: ComBat PCA diagnostics", "S5: Bootstrap LASSO stability", "S6: Nomogram sensitivity diagnostics",
            "S7: Gene-level ablation forest plot", "S8: TGF-beta pathway gene set composition"]:
    add_para(f"  {fig}", size=9)

# ====== SAVE ======
doc.save(OUT_PATH)
print(f"Supplementary DOCX saved: {OUT_PATH}")
print(f"Size: {os.path.getsize(OUT_PATH)/1024:.0f} KB")
