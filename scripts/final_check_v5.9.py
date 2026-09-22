"""
Final check: verify no remaining duplicate content.
"""
from docx import Document

SRC = "/path/to/xelox_project/reports/manuscript_draft_v5.9_submission.docx"
doc = Document(SRC)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check tables for duplicates
print("\n=== Table fingerprint comparison ===")
table_fingerprints = []
for i, table in enumerate(doc.tables):
    if len(table.rows) > 0:
        fp = ' | '.join(cell.text.strip()[:25] for cell in table.rows[0].cells)
        table_fingerprints.append((i, fp))

# Group by fingerprint
from collections import defaultdict
fp_groups = defaultdict(list)
for idx, fp in table_fingerprints:
    fp_groups[fp].append(idx)

print("Duplicate table groups:")
found_dup = False
for fp, indices in fp_groups.items():
    if len(indices) > 1:
        found_dup = True
        print(f"  DUPLICATE: Tables {indices} — {fp}")

if not found_dup:
    print("  No duplicate tables found ✓")

# Check for "File Organization", "Document Summary", etc.
print("\n=== Checking for removed content ===")
bad_markers = ["File Organization", "Resource Summary", "Document Summary",
               "References (GEO Datasets)", "Directory Structure",
               "Supplementary Tables S1-S16 and Figures S1-S12 are provided in a separate file"]

for marker in bad_markers:
    for i, p in enumerate(doc.paragraphs):
        if marker in p.text:
            print(f"  WARNING: Found '{marker}' at Para {i}")

print("\n=== Block structure ===")
supp_count = 0
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t == "Supplementary Materials":
        supp_count += 1
        # Show next few paragraphs
        print(f"  Block {supp_count} (Para {i}):")
        for j in range(1, 4):
            if i+j < len(doc.paragraphs):
                next_t = doc.paragraphs[i+j].text.strip()
                if next_t:
                    print(f"    +{j}: {next_t[:80]}")

# Verify Table S33 correctness
print("\n=== Table S33 data verification ===")
for i, table in enumerate(doc.tables):
    first_row = ' '.join(cell.text.strip() for cell in table.rows[0].cells)
    if "Metric" in first_row and "Value" in first_row:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any("Bootstrap" in c for c in cells) and any("C-index" in c for c in cells):
                print(f"  Table {i}: {cells}")
            if any("Optimism" in c for c in cells):
                print(f"  Table {i}: {cells}")
